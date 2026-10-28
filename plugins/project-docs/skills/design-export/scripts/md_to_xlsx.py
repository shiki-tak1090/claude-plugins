# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl==3.1.5", "pillow==13.0.0"]
# [tool.uv]
# exclude-newer = "2026-10-28T00:00:00Z"
# ///
"""設計書 Markdown を Excel（xlsx）に変換する。

章の見出しごとにシートを分け、見出し・本文・表を上から順に書き出す。章の見出しは、
文書内で 2 回以上使われている最も上位のレベルとする（# を章に使う文書と ## を章に使う文書の両方に対応するため）。
すべての値は文字列として書き、先頭の 0 や "=" で始まる値が変わらないようにする。
テンプレートの xlsx を指定した場合は、その既存シート（表紙など）の後ろにシートを追加する。
"""
from __future__ import annotations

import argparse
import re
import sys
import tempfile
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(Path(__file__).parent))
from mdmodel import (  # noqa: E402
    Block, expand_sources, find_mmdc, output_path, parse, plain, render_mermaid, resolve_embed,
)

THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
HEADING_SIZE = {1: 16, 2: 14, 3: 12, 4: 11}


def display_width(s: str) -> int:
    return max((sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in line)
                for line in s.splitlines() or [""]), default=0)


def sheet_title(text: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "", plain(text)).strip()[:28] or "Sheet"
    title, n = base, 2
    while title in used:
        title, n = f"{base[:25]}_{n}", n + 1
    used.add(title)
    return title


def put(ws, row: int, col: int, value: str, **style):
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # 数式として解釈させない
    for k, v in style.items():
        setattr(cell, k, v)
    return cell


def split_sections(blocks: list[Block]) -> list[tuple[str, list[Block]]]:
    levels = [b.level for b in blocks if b.kind == "heading"]
    split_level = min((lv for lv in set(levels) if levels.count(lv) >= 2), default=2)
    sections: list[tuple[str, list[Block]]] = [("概要", [])]
    for b in blocks:
        if b.kind == "heading" and b.level == split_level:
            sections.append((b.text, []))
        sections[-1][1].append(b)
    return [s for s in sections if s[1]]


def write_section(ws, blocks: list[Block], tmpdir: Path, diagram_no: list[int]) -> None:
    widths: dict[int, int] = {}
    row = 1
    for b in blocks:
        if b.kind == "heading":
            put(ws, row, 1, plain(b.text), font=Font(bold=True, size=HEADING_SIZE.get(b.level, 11)))
            row += 1
        elif b.kind == "paragraph":
            put(ws, row, 1, plain(b.text))
            row += 1
        elif b.kind == "list_item":
            mark = {"todo": "☐ ", "done": "☑ "}.get(b.lang, "・")
            put(ws, row, 1, "  " * b.level + mark + plain(b.text))
            row += 1
        elif b.kind == "code":
            put(ws, row, 1, b.text, font=Font(name="Consolas"), alignment=Alignment(wrap_text=True, vertical="top"))
            row += 1
        elif b.kind == "table":
            for r_i, r in enumerate(b.rows):
                for c_i, value in enumerate(r, start=1):
                    v = plain(value)
                    cell = put(ws, row, c_i, v, border=BORDER,
                               alignment=Alignment(wrap_text=True, vertical="top"))
                    if r_i == 0:
                        cell.font = Font(bold=True)
                        cell.fill = HEADER_FILL
                    widths[c_i] = max(widths.get(c_i, 0), min(display_width(v), 60))
                row += 1
        elif b.kind == "mermaid":
            diagram_no[0] += 1
            if find_mmdc():
                from openpyxl.drawing.image import Image
                png = tmpdir / f"diagram{diagram_no[0]}.png"
                render_mermaid(b.text, png)
                img = Image(str(png))
                ws.add_image(img, f"A{row}")
                row += max(1, int(img.height / 20))
            else:
                put(ws, row, 1, "[図] mmdc がないため Mermaid のソースを記載", font=Font(italic=True))
                put(ws, row + 1, 1, b.text, alignment=Alignment(wrap_text=True, vertical="top"))
                row += 2
        elif b.kind == "image":
            if b.lang:
                from openpyxl.drawing.image import Image
                img = Image(b.lang)
                ws.add_image(img, f"A{row}")
                row += max(1, int(img.height / 20))
            else:
                put(ws, row, 1, f"[画像が見つかりません: {b.text}]", font=Font(italic=True))
                row += 1
        row += 1  # ブロック間の空行
    for col, w in widths.items():
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = max(w + 2, 8)


def convert(src: Path, out: Path, template: Path | None) -> None:
    blocks = parse(src.read_text(encoding="utf-8"))
    for b in blocks:
        if b.kind == "image":
            path = resolve_embed(b.lang, src.parent)
            b.text, b.lang = b.lang, (str(path) if path else "")
    if template:
        wb = load_workbook(template)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    used = set(wb.sheetnames)
    diagram_no = [0]
    with tempfile.TemporaryDirectory() as tmp:
        for title, section in split_sections(blocks):
            write_section(wb.create_sheet(sheet_title(title, used)), section, Path(tmp), diagram_no)
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("src", nargs="+", help="変換する Markdown ファイルまたはフォルダ")
    ap.add_argument("--src-root", default=".", help="出力先の相対パスを決める基準")
    ap.add_argument("--out-root", default="exports/xlsx")
    ap.add_argument("--template", default="99_テンプレート/export/template.xlsx")
    args = ap.parse_args()

    template = Path(args.template)
    if not template.exists():
        template = None
    for src in expand_sources(args.src):
        out = output_path(src, Path(args.src_root), Path(args.out_root), ".xlsx")
        convert(src, out, template)
        print(f"{src} -> {out}")


if __name__ == "__main__":
    main()
